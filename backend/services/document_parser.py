"""
backend/services/document_parser.py

Extracts text from PDF/DOCX tender documents.
Preserves page/section information for evidence tracking.
"""

import io
from typing import List, Dict


def extract_pdf_text(file_bytes: bytes) -> List[Dict]:
    """
    Extract text from PDF using pdfplumber.
    Returns list of {page, text} dicts.
    """
    pages = []
    try:
        import pdfplumber
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for i, page in enumerate(pdf.pages, start=1):
                text = page.extract_text() or ""
                if text.strip():
                    pages.append({"page": i, "text": text.strip()})
    except Exception as e:
        print(f"[DocumentParser] PDF extraction error: {e}")
    return pages


def extract_docx_text(file_bytes: bytes) -> List[Dict]:
    """
    Extract text from DOCX using python-docx.
    Returns list of {page, text} dicts.
    """
    pages = []
    try:
        from docx import Document
        doc = Document(io.BytesIO(file_bytes))
        full_text = []
        for para in doc.paragraphs:
            if para.text.strip():
                full_text.append(para.text.strip())
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                if row_text:
                    full_text.append(row_text)
        if full_text:
            pages.append({"page": 1, "text": "\n".join(full_text)})
    except Exception as e:
        print(f"[DocumentParser] DOCX extraction error: {e}")
    return pages


def extract_text(file_bytes: bytes, filename: str) -> List[Dict]:
    """
    Auto-detect file type and extract text.
    Returns list of {page, text} dicts.
    """
    fname = filename.lower()
    if fname.endswith(".pdf"):
        pages = extract_pdf_text(file_bytes)
        if not pages:
            print("[DocumentParser] PDF text extraction empty — document may be scanned.")
        return pages
    elif fname.endswith(".docx"):
        return extract_docx_text(file_bytes)
    else:
        print(f"[DocumentParser] Unsupported file type: {filename}")
        return []


def combine_pages_text(pages: List[Dict], max_chars: int = 10000) -> str:
    """Combine page texts into a single string, truncated."""
    combined = "\n\n".join(
        f"[Page {p['page']}]\n{p['text']}" for p in pages
    )
    return combined[:max_chars]
